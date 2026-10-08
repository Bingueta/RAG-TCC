"""Parte 4: a planilha Excel de cada execução (formato na seção 3.4 de docs/divisao-tarefas.md).

Três abas:
- respostas: exatamente as 6 colunas pedidas, nesta ordem: id, titulo, tematica_1,
  tematica_2, metodologias (separadas por "; ") e modelo;
- detalhes: id, evidência de cada campo como TEXTO das frases (quem lê a planilha não
  precisa procurar F2 no resumo), status, técnica e versão do prompt;
- execucao: a máquina, a configuração e a data da rodada.

Célula vazia na aba respostas sempre tem o motivo na coluna status da aba detalhes.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.contratos import Campo, Dissertacao, DissertacaoUnitarizada, RespostaIA

COLUNAS_RESPOSTAS = ["id", "titulo", "tematica_1", "tematica_2", "metodologias", "modelo"]
COLUNAS_DETALHES = ["id", "evidencia_tematica_1", "evidencia_tematica_2", "evidencia_metodologias",
                    "status", "tecnica", "versao_prompt"]
SEPARADOR_METODOLOGIAS = "; "
NOME_DO_ARQUIVO = "respostas.xlsx"

# Largura das colunas (em caracteres), para a planilha abrir legível sem ajuste manual.
_LARGURAS = {"id": 7, "titulo": 50, "tematica_1": 34, "tematica_2": 34, "metodologias": 60, "modelo": 22,
             "evidencia_tematica_1": 60, "evidencia_tematica_2": 60, "evidencia_metodologias": 80,
             "status": 30, "tecnica": 10, "versao_prompt": 8, "campo": 28, "valor": 90}
_CABECALHO = Font(bold=True, color="FFFFFF")
_FUNDO_CABECALHO = PatternFill("solid", fgColor="305496")


def texto_das_metodologias(resposta: RespostaIA) -> str:
    return SEPARADOR_METODOLOGIAS.join(m.texto for m in resposta.metodologias if m.texto)


def status_da_resposta(resposta: RespostaIA) -> str:
    """"ok", ou o que não está ok, campo por campo (ex.: "tematica_2: sem_evidencia")."""
    campos = [resposta.tematica_1, resposta.tematica_2, *resposta.metodologias]
    if all(c.status == "erro" for c in campos):
        return "erro"
    problemas = [f"{nome}: {campo.status}" for nome, campo in
                 (("tematica_1", resposta.tematica_1), ("tematica_2", resposta.tematica_2)) if campo.status != "ok"]
    for m in resposta.metodologias:
        if m.status == "nao_informado":
            problemas.append("metodologias: nao_informado")
        elif m.status != "ok":
            problemas.append(f"metodologia '{m.texto}': {m.status}")
    return "; ".join(problemas) or "ok"


def _frases_citadas(campos: list[Campo], frases: dict[str, str]) -> str:
    """Cada frase citada uma vez, na ordem do resumo: "F3: Trata-se de…"."""
    ids = sorted({i for c in campos for i in c.evidencia}, key=lambda i: int(i[1:]))
    return "\n".join(f"{i}: {frases.get(i, '(frase não encontrada no resumo)')}" for i in ids)


def _evidencia_das_metodologias(resposta: RespostaIA, frases: dict[str, str]) -> str:
    """Os métodos com os códigos e, embaixo, cada frase citada uma vez.

    Vários métodos costumam vir da mesma frase ("pesquisa qualitativa, estudo de caso"):
    repetir a frase em cada um deixaria a célula enorme sem dizer nada a mais.
    """
    com_evidencia = [m for m in resposta.metodologias if m.evidencia]
    if not com_evidencia:
        return ""
    metodos = SEPARADOR_METODOLOGIAS.join(f"{m.texto} [{', '.join(m.evidencia)}]" for m in com_evidencia)
    return metodos + "\n\n" + _frases_citadas(com_evidencia, frases)


def _unitarizada(item) -> DissertacaoUnitarizada:
    if isinstance(item, DissertacaoUnitarizada):
        return item
    from src.unitarizar import unitarizar  # spaCy: só quando o corpus vem sem as frases

    return unitarizar(item)


def _linhas_da_execucao(execucao: dict) -> list[tuple[str, str]]:
    """O execucao.json em pares campo/valor, com os dicionários abertos e um resumo do tempo."""
    linhas = []
    for campo, valor in execucao.items():
        if campo in ("segundos", "tentativas"):
            continue
        if isinstance(valor, dict):
            linhas += [(f"{campo}.{k}", str(v)) for k, v in valor.items()]
        else:
            linhas.append((campo, str(valor)))
    segundos = list(execucao.get("segundos", {}).values())
    if segundos:
        linhas.append(("tempo_medio_por_dissertacao_s", f"{sum(segundos) / len(segundos):.1f}"))
        linhas.append(("tempo_total_de_geracao_s", f"{sum(segundos):.0f}"))
    repetidas = sum(1 for t in execucao.get("tentativas", {}).values() if t > 1)
    if execucao.get("tentativas"):
        linhas.append(("dissertacoes_que_precisaram_de_2a_tentativa", str(repetidas)))
    return linhas


def _aba(planilha, titulo: str, colunas: list[str], linhas: list[list], nova: bool = True):
    aba = planilha.create_sheet(titulo) if nova else planilha.active
    aba.title = titulo
    aba.append(colunas)
    for linha in linhas:
        aba.append(linha)
    for numero, coluna in enumerate(colunas, start=1):
        celula = aba.cell(row=1, column=numero)
        celula.font, celula.fill = _CABECALHO, _FUNDO_CABECALHO
        aba.column_dimensions[get_column_letter(numero)].width = _LARGURAS.get(coluna, 20)
    for linha in aba.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(wrap_text=True, vertical="top")
    aba.freeze_panes = "B2"
    aba.auto_filter.ref = aba.dimensions
    return aba


def exportar_excel(respostas: list[RespostaIA], corpus: list[Dissertacao | DissertacaoUnitarizada],
                   pasta: str | Path, execucao: dict | None = None) -> Path:
    """Grava <pasta>/respostas.xlsx, uma linha por dissertação do corpus, na ordem dele.

    corpus pode vir com as frases (DissertacaoUnitarizada) ou sem (Dissertacao; aí o resumo
    é dividido de novo, com o mesmo resultado). Dissertação sem resposta ainda (execução
    interrompida) sai com as células vazias e status "sem resposta".
    """
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    por_id = {r.dissertacao_id: r for r in respostas}
    linhas_respostas, linhas_detalhes = [], []
    for item in corpus:
        unitarizada = _unitarizada(item)
        d = unitarizada.dissertacao
        frases = {f.id: f.texto for f in unitarizada.frases}
        r = por_id.get(d.id)
        if r is None:
            linhas_respostas.append([d.id, d.titulo, "", "", "", ""])
            linhas_detalhes.append([d.id, "", "", "", "sem resposta", "", ""])
            continue
        linhas_respostas.append([d.id, d.titulo, r.tematica_1.texto, r.tematica_2.texto,
                                 texto_das_metodologias(r), r.modelo])
        linhas_detalhes.append([d.id, _frases_citadas([r.tematica_1], frases), _frases_citadas([r.tematica_2], frases),
                                _evidencia_das_metodologias(r, frases), status_da_resposta(r), r.tecnica,
                                r.versao_prompt])

    planilha = Workbook()
    _aba(planilha, "respostas", COLUNAS_RESPOSTAS, linhas_respostas, nova=False)
    _aba(planilha, "detalhes", COLUNAS_DETALHES, linhas_detalhes)
    _aba(planilha, "execucao", ["campo", "valor"], [list(l) for l in _linhas_da_execucao(execucao or {})])
    caminho = pasta / NOME_DO_ARQUIVO
    planilha.save(caminho)
    return caminho
