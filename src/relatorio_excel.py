"""Parte 5: relatório em Excel (tarefa 6), a partir dos JSON que o comparar já gravou.

- comparacao.xlsx: aba "execucoes" (uma linha por execução, com modelo, técnica, prompt e
  os indicadores lado a lado), "pareadas" (diferenças pareadas, quando há referência),
  "concordancia" (pares de execuções e cada dissertação) e "graficos", com um gráfico de
  barras por indicador.
- relatorio.xlsx de cada execução: aba "resumo" (indicadores sem gabarito e, se houver,
  contra a referência), "configuracao" (o execucao.json) e "por_dissertacao" (só com
  referência; esse arquivo lista os termos esperados e fica em data/avaliacao/, fora do git).

Lê os JSON em vez de recalcular. Assim o Excel mostra exatamente o que foi medido, e o
src/comparar.py, que é a régua, não precisa mudar (o sha256 dele é o que prova que as
rodadas foram medidas do mesmo jeito).

Comando: python -m src.relatorio_excel PASTA
    converte todo comparacao.json e relatorio.json abaixo de PASTA, ao lado de cada um.
"""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Indicadores que ganham gráfico, na ordem em que aparecem na aba "graficos". Só entra
# o que existir na tabela: sem referência, os de metodologia e temática não existem.
INDICADORES_GRAFICO = [
    ("met_f1", "Metodologias: F1"),
    ("met_precisao", "Metodologias: precisão"),
    ("met_revocacao", "Metodologias: revocação"),
    ("tem_acerto_1", "Temáticas: acerto@1"),
    ("tem_margem", "Temáticas: margem sobre a linha de base"),
    ("itens_exatos", "Formato: itens com nome limpo"),
    ("eixos_cobertos", "Eixos de metodologia cobertos (de 4)"),
    ("campos_ok", "Campos com status ok"),
    ("tempo_medio_s", "Tempo médio por dissertação (s)"),
    ("palavras_por_tematica", "Palavras por temática"),
]


def _celula(valor):
    """Valor que o Excel aceita numa célula: lista vira texto com "; ", dicionário vira JSON."""
    if isinstance(valor, (list, tuple)):
        return "; ".join(str(v) for v in valor)
    if isinstance(valor, dict):
        return json.dumps(valor, ensure_ascii=False)
    return valor


def _achatar(dados: dict, prefixo: str = "") -> dict:
    """{"status": {"por_campo": {"ok": 3}}} → {"status.por_campo.ok": 3}."""
    plano = {}
    for chave, valor in dados.items():
        nome = f"{prefixo}{chave}"
        if isinstance(valor, dict):
            plano |= _achatar(valor, f"{nome}.")
        else:
            plano[nome] = valor
    return plano


def _aba(livro: Workbook, titulo: str, linhas: list[dict], primeira: bool = False):
    """Uma aba com cabeçalho em negrito e uma linha por dicionário; as colunas são a união das chaves."""
    aba = livro.active if primeira else livro.create_sheet()
    aba.title = titulo
    colunas = list(dict.fromkeys(k for l in linhas for k in l))
    aba.append(colunas)
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for l in linhas:
        aba.append([_celula(l.get(c)) for c in colunas])
    for i, coluna in enumerate(colunas, start=1):
        largura = max([len(str(coluna))] + [len(str(_celula(l.get(coluna)) or "")) for l in linhas])
        aba.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 8), 60)
        for celula in aba[get_column_letter(i)][1:]:
            if isinstance(celula.value, float):
                celula.number_format = "0.000"
    aba.freeze_panes = "B2"
    return aba, colunas


def _graficos(livro: Workbook, aba_execucoes, colunas: list[str], n: int) -> int:
    """Um gráfico de barras por indicador, com as execuções no eixo. Devolve quantos fez."""
    aba = livro.create_sheet("graficos")
    feitos = 0
    categorias = Reference(aba_execucoes, min_col=1, min_row=2, max_row=n + 1)
    for chave, titulo in INDICADORES_GRAFICO:
        if chave not in colunas:
            continue
        coluna = colunas.index(chave) + 1
        valores = [aba_execucoes.cell(row=r, column=coluna).value for r in range(2, n + 2)]
        if not any(isinstance(v, (int, float)) for v in valores):
            continue
        grafico = BarChart()
        grafico.type = "bar"           # barras deitadas: os nomes das execuções são longos
        grafico.title = titulo
        grafico.legend = None
        grafico.height, grafico.width = 7 + 0.35 * n, 16
        grafico.add_data(Reference(aba_execucoes, min_col=coluna, min_row=1, max_row=n + 1), titles_from_data=True)
        grafico.set_categories(categorias)
        aba.add_chart(grafico, f"{'A' if feitos % 2 == 0 else 'K'}{1 + (feitos // 2) * int(16 + 0.7 * n)}")
        feitos += 1
    return feitos


def gravar_comparacao(comparacao: dict, caminho: str | Path) -> Path:
    """O comparacao.json em Excel: tabela das execuções, pareadas, concordância e gráficos."""
    livro = Workbook()
    execucoes = comparacao.get("execucoes") or []
    aba_execucoes, colunas = _aba(livro, "execucoes", execucoes, primeira=True)
    pares = comparacao.get("diferencas_pareadas") or []
    if pares:
        _aba(livro, "pareadas", [{"muda": p["muda"], "a": p["a"], "b": p["b"],
                                  "f1_media": p["f1"]["media"], "f1_ic95_baixo": (p["f1"]["ic95"] or [None])[0],
                                  "f1_ic95_alto": (p["f1"]["ic95"] or [None, None])[1],
                                  "ganhou": p["f1"]["ganhou"], "perdeu": p["f1"]["perdeu"],
                                  "empatou": p["f1"]["empatou"],
                                  "tematicas_media": (p.get("similaridade_tematicas") or {}).get("media")}
                                 for p in pares])
    concordancia = comparacao.get("concordancia")
    if concordancia:
        _aba(livro, "concordancia", concordancia["pares"])
        _aba(livro, "concordancia_por_dissertacao",
             [{"dissertacao_id": d, **v} for d, v in concordancia["por_dissertacao"].items()])
    if execucoes:
        _graficos(livro, aba_execucoes, colunas, len(execucoes))
    caminho = Path(caminho)
    livro.save(caminho)
    return caminho


def gravar_relatorio(relatorio: dict, caminho: str | Path) -> Path:
    """O relatorio.json de uma execução em Excel."""
    livro = Workbook()
    resumo = {k: relatorio.get(k) for k in ("execucao", "modelo", "tecnica", "versao_prompt")}
    resumo |= _achatar(relatorio.get("sem_gabarito") or {}, "sem_gabarito.")
    contra = relatorio.get("contra_referencia")
    if contra:
        resumo |= _achatar({k: v for k, v in contra.items() if k != "por_dissertacao"}, "contra_referencia.")
    _aba(livro, "resumo", [{"indicador": k, "valor": _celula(v)} for k, v in resumo.items()], primeira=True)
    if relatorio.get("configuracao"):
        _aba(livro, "configuracao", [{"campo": k, "valor": _celula(v)}
                                     for k, v in _achatar(relatorio["configuracao"]).items()])
    if contra and contra.get("por_dissertacao"):
        _aba(livro, "por_dissertacao", contra["por_dissertacao"])
    caminho = Path(caminho)
    livro.save(caminho)
    return caminho


def converter_pasta(pasta: str | Path) -> list[Path]:
    """Grava o .xlsx ao lado de cada comparacao.json e relatorio.json abaixo de pasta."""
    feitos = []
    for arquivo in sorted(Path(pasta).rglob("*.json")):
        if arquivo.name == "comparacao.json":
            feitos.append(gravar_comparacao(json.loads(arquivo.read_text(encoding="utf-8")),
                                            arquivo.with_suffix(".xlsx")))
        elif arquivo.name == "relatorio.json":
            feitos.append(gravar_relatorio(json.loads(arquivo.read_text(encoding="utf-8")),
                                           arquivo.with_suffix(".xlsx")))
    return feitos


def main() -> None:
    """Comando: python -m src.relatorio_excel PASTA"""
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pasta")
    for caminho in converter_pasta(parser.parse_args().pasta):
        print(caminho)


if __name__ == "__main__":
    main()
