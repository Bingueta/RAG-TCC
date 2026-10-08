"""Parte 4: testes do pipeline, em modo mock (sem Ollama, sem índice e sem spaCy)."""

import hashlib
import json

import pytest
from openpyxl import load_workbook

from src import mocks, pipeline
from src.exportar import COLUNAS_DETALHES, COLUNAS_RESPOSTAS
from src.pipeline import CorpusDesconhecido, executar, nome_da_pasta, tentativas
from src.sugerir import SEPARADOR_TENTATIVAS


def rodar(tmp_path, tecnica="hibrido", **extra):
    return executar("mock", tecnica, mock=True, base=tmp_path, **extra)


def sha(caminho):
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def test_nome_da_pasta_serve_no_windows():
    assert nome_da_pasta("qwen3.5:9b-q4_K_M", "hibrido", "v3") == "qwen3.5-9b-q4_K_M__hibrido__v3"
    assert nome_da_pasta("a/b:c", "denso", "v1", "busca-corrigida") == "a-b-c__denso__v1__busca-corrigida"


def test_mock_grava_as_tres_saidas(tmp_path):
    pasta = rodar(tmp_path)
    assert pasta == tmp_path / "mock" / "mock__hibrido__v3"
    respostas = json.loads((pasta / "respostas.json").read_text(encoding="utf-8"))
    assert [r["dissertacao_id"] for r in respostas] == ["E001", "E002", "E003"]
    execucao = json.loads((pasta / "execucao.json").read_text(encoding="utf-8"))
    assert execucao["corpus"] == "mock" and execucao["versao_prompt"] == "v3" and execucao["erros"] == 0
    assert set(execucao["segundos"]) == set(execucao["tentativas"]) == {"E001", "E002", "E003"}
    assert "openpyxl" in execucao["bibliotecas"]
    assert (pasta / "respostas.xlsx").exists()


def test_planilha_tem_as_6_colunas_na_ordem_e_metodologias_com_ponto_e_virgula(tmp_path):
    planilha = load_workbook(rodar(tmp_path) / "respostas.xlsx")
    assert planilha.sheetnames == ["respostas", "detalhes", "execucao"]
    linhas = list(planilha["respostas"].iter_rows(values_only=True))
    assert list(linhas[0]) == COLUNAS_RESPOSTAS
    assert len(linhas) == 4  # cabeçalho + 3 dissertações
    e001 = linhas[1]
    assert e001[4] == "Pesquisa qualitativa; Estudo de caso; Entrevista semiestruturada; Análise de conteúdo"
    assert all(linha[5] == "mock" for linha in linhas[1:])
    assert linhas[3][4] == "Não informado no resumo"


def test_detalhes_mostram_a_evidencia_como_texto_da_frase(tmp_path):
    planilha = load_workbook(rodar(tmp_path) / "respostas.xlsx")
    linhas = list(planilha["detalhes"].iter_rows(values_only=True))
    assert list(linhas[0]) == COLUNAS_DETALHES
    e001 = dict(zip(COLUNAS_DETALHES, linhas[1]))
    assert e001["evidencia_tematica_1"].startswith("F1: Esta pesquisa investiga")
    assert "Estudo de caso [F3]" in e001["evidencia_metodologias"]
    # a frase F3 justifica dois métodos, mas aparece uma vez só
    assert e001["evidencia_metodologias"].count("F3: Trata-se de uma pesquisa qualitativa") == 1
    assert e001["status"] == "ok" and e001["tecnica"] == "hibrido" and e001["versao_prompt"] == "v3"
    e003 = dict(zip(COLUNAS_DETALHES, linhas[3]))
    assert e003["status"] == "metodologias: nao_informado"


def test_aba_execucao_registra_a_rodada(tmp_path):
    planilha = load_workbook(rodar(tmp_path) / "respostas.xlsx")
    campos = {linha[0]: linha[1] for linha in planilha["execucao"].iter_rows(min_row=2, values_only=True)}
    assert campos["modelo"] == "mock" and campos["tecnica"] == "hibrido" and campos["versao_prompt"] == "v3"
    assert "maquina.cpu" in campos and "bibliotecas.openpyxl" in campos and "tempo_medio_por_dissertacao_s" in campos


def test_retoma_de_onde_parou(tmp_path, monkeypatch):
    pasta = rodar(tmp_path)
    respostas = json.loads((pasta / "respostas.json").read_text(encoding="utf-8"))
    (pasta / "respostas.json").write_text(json.dumps(respostas[:1], ensure_ascii=False), encoding="utf-8")
    chamadas = []
    original = mocks.gerar_resposta
    monkeypatch.setattr(mocks, "gerar_resposta", lambda u, *a, **k: chamadas.append(u.dissertacao.id) or original(u, *a, **k))
    rodar(tmp_path)
    assert chamadas == ["E002", "E003"]  # E001 já estava pronta
    execucao = json.loads((pasta / "execucao.json").read_text(encoding="utf-8"))
    assert "retomada_em" in execucao and set(execucao["segundos"]) == {"E001", "E002", "E003"}


def test_execucao_completa_nao_e_regravada(tmp_path):
    pasta = rodar(tmp_path)
    antes = (sha(pasta / "respostas.json"), sha(pasta / "execucao.json"))
    (pasta / "respostas.xlsx").unlink()
    rodar(tmp_path)
    assert (sha(pasta / "respostas.json"), sha(pasta / "execucao.json")) == antes
    assert (pasta / "respostas.xlsx").exists()  # só a planilha é feita de novo


def test_sobrescrever_refaz_tudo(tmp_path, monkeypatch):
    rodar(tmp_path)
    chamadas = []
    original = mocks.gerar_resposta
    monkeypatch.setattr(mocks, "gerar_resposta", lambda u, *a, **k: chamadas.append(u.dissertacao.id) or original(u, *a, **k))
    rodar(tmp_path, sobrescrever=True)
    assert chamadas == ["E001", "E002", "E003"]


def test_erro_numa_dissertacao_nao_para_as_outras(tmp_path, monkeypatch):
    original = mocks.recuperar

    def recuperar_que_falha(u, tecnica):
        if u.dissertacao.id == "E002":
            raise RuntimeError("índice corrompido")
        return original(u, tecnica)

    monkeypatch.setattr(mocks, "recuperar", recuperar_que_falha)
    pasta = rodar(tmp_path)
    respostas = {r["dissertacao_id"]: r for r in json.loads((pasta / "respostas.json").read_text(encoding="utf-8"))}
    assert respostas["E002"]["tematica_1"]["status"] == "erro"
    assert "RuntimeError: índice corrompido" in respostas["E002"]["saida_bruta"]
    assert respostas["E001"]["tematica_1"]["status"] == respostas["E003"]["tematica_1"]["status"] == "ok"
    assert json.loads((pasta / "execucao.json").read_text(encoding="utf-8"))["erros"] == 1
    status = [linha[4] for linha in load_workbook(pasta / "respostas.xlsx")["detalhes"].iter_rows(min_row=2, values_only=True)]
    assert status == ["ok", "erro", "metodologias: nao_informado"]


def test_sem_rag_do_mock_manda_todas_as_frases(tmp_path):
    respostas = json.loads((rodar(tmp_path, tecnica="sem_rag") / "respostas.json").read_text(encoding="utf-8"))
    assert all(r["tecnica"] == "sem_rag" for r in respostas)


def test_ids_rodam_so_algumas(tmp_path):
    respostas = json.loads((rodar(tmp_path, ids=["E003"]) / "respostas.json").read_text(encoding="utf-8"))
    assert [r["dissertacao_id"] for r in respostas] == ["E003"]


def test_corpus_desconhecido_e_erro_claro(tmp_path):
    with pytest.raises(CorpusDesconhecido, match="corpus 'as59' não existe"):
        executar("mock", "hibrido", corpus="as59", mock=True, base=tmp_path)


def test_calibracao_vai_para_a_subpasta():
    assert pipeline.pasta_da_execucao("m:1b", "denso", "v3", "calibracao").parent == pipeline.PASTA_SUGESTOES / "calibracao"
    assert pipeline.pasta_da_execucao("m:1b", "denso", "v3", "corpus").parent == pipeline.PASTA_SUGESTOES


def test_tentativas_conta_so_as_respostas_do_modelo():
    from src.contratos import Campo, RespostaIA

    def com_saida(saida):
        c = Campo(texto="x", evidencia=["F1"])
        return RespostaIA("E001", c, c, [c], "m", "v3", "hibrido", saida_bruta=saida)

    assert tentativas(com_saida("{}")) == 1
    assert tentativas(com_saida(SEPARADOR_TENTATIVAS.join(["{quebrado", "{}"]))) == 2
    assert tentativas(com_saida(SEPARADOR_TENTATIVAS.join(["{quebrado", "{quebrado", "[não lida: JSON quebrado]"]))) == 2
    assert tentativas(com_saida("[falha: ConnectionError: recusada]")) == 0


def test_main_em_modo_mock(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(pipeline, "PASTA_SUGESTOES", tmp_path)
    pipeline.main(["--mock", "--todos"])
    assert sorted(p.name for p in (tmp_path / "mock").iterdir()) == [
        "mock__denso__v3", "mock__hibrido__v3", "mock__sem_rag__v3"]
    assert "planilha:" in capsys.readouterr().out
